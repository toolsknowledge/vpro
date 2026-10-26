"""Shared test fixtures: sample files built in memory and a fake OpenAI client.

No test in this folder (except tests/integration) calls the real OpenAI API.
"""

from __future__ import annotations

import hashlib
import io
import math
import re
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from backend.document_registry import DocumentRegistry
from backend.embeddings import EmbeddingService
from backend.indexing_pipeline import IndexingPipeline
from backend.vector_store import VectorStore

FAKE_DIM = 64


def fake_vector(text: str, dim: int = FAKE_DIM) -> list[float]:
    """Deterministic bag-of-words vector: texts sharing words get similar vectors."""
    vec = [0.0] * dim
    for word in re.findall(r"[a-z0-9]+", text.lower()):
        bucket = int(hashlib.md5(word.encode()).hexdigest(), 16) % dim
        vec[bucket] += 1.0
    norm = math.sqrt(sum(v * v for v in vec)) or 1.0
    return [v / norm for v in vec]


class FakeEmbeddings:
    def __init__(self, dim=FAKE_DIM, fail_times: int = 0, error: Exception | None = None):
        self.dim, self.fail_times, self.error, self.calls = dim, fail_times, error, 0

    def create(self, model, input):
        self.calls += 1
        if self.fail_times > 0:
            self.fail_times -= 1
            raise self.error
        data = [SimpleNamespace(index=i, embedding=fake_vector(t, self.dim)) for i, t in enumerate(input)]
        return SimpleNamespace(data=list(reversed(data)))  # reversed: the code must sort by index


class FakeChat:
    def __init__(self, reply: str = "Answer [S1]."):
        self.reply, self.last_kwargs = reply, None
        self.completions = self

    def create(self, **kwargs):
        self.last_kwargs = kwargs
        message = SimpleNamespace(content=self.reply)
        usage = SimpleNamespace(prompt_tokens=10, completion_tokens=5)
        return SimpleNamespace(choices=[SimpleNamespace(message=message)], usage=usage)


class FakeClient:
    def __init__(self, reply: str = "Answer [S1].", **embedding_kwargs):
        self.embeddings = FakeEmbeddings(**embedding_kwargs)
        self.chat = FakeChat(reply)


# --------------------------------------------------------------- file builders
def make_pdf_bytes(pages: list[str]) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    for text in pages:
        y = 800
        for line in text.split("\n"):
            if line:
                pdf.drawString(50, y, line)
            y -= 16
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_xlsx_bytes(sheets: dict[str, list[tuple]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def make_docx_bytes(blocks: list[tuple[str, str]], table: list[tuple] | None = None) -> bytes:
    """blocks = [("heading", text) | ("para", text)]"""
    doc = Document()
    for kind, text in blocks:
        if kind == "heading":
            doc.add_heading(text, 1)
        else:
            doc.add_paragraph(text)
    if table:
        t = doc.add_table(rows=0, cols=len(table[0]))
        for row in table:
            cells = t.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = str(value)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def course_xlsx() -> bytes:
    return make_xlsx_bytes(
        {
            "Course Catalog": [
                ("Course", "Duration", "Fee"),
                ("Python Full Stack", "16 weeks", 30000),
                ("Data Science", "20 weeks", None),
            ],
            "Batches": [("Course", "Start"), ("Python Full Stack", "2026-11-10")],
        }
    )


@pytest.fixture
def services_factory(tmp_path):
    """Build a real registry + real ChromaDB (in tmp_path) + fake OpenAI embeddings."""

    def build(client=None, model="fake-embedding-model", collection="test_collection"):
        client = client or FakeClient()
        registry = DocumentRegistry(tmp_path / "registry.json", tmp_path / "traces", tmp_path / "uploads")
        store = VectorStore(tmp_path / "chroma", collection)
        embedder = EmbeddingService(client, model, batch_size=4, max_attempts=3)
        pipeline = IndexingPipeline(registry, store, embedder, max_upload_bytes=5 * 1024 * 1024)
        return SimpleNamespace(client=client, registry=registry, store=store, embedder=embedder, pipeline=pipeline)

    return build


@pytest.fixture(autouse=True)
def no_sleep(monkeypatch):
    """Retries back off with time.sleep; skip the waiting in tests."""
    monkeypatch.setattr("backend.openai_client.time.sleep", lambda seconds: None)
