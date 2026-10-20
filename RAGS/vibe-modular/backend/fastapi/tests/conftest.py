from __future__ import annotations

import io
import math
import re
import zlib
from collections.abc import Callable, Iterator
from types import SimpleNamespace
from typing import Any

import pytest
from docx import Document

from app.core.config import Settings

FAKE_DIMENSION = 32


@pytest.fixture(autouse=True)
def _isolate_environment(monkeypatch: pytest.MonkeyPatch) -> None:
    """Keep the developer's shell environment from leaking into Settings during tests."""
    for name in Settings.model_fields:
        monkeypatch.delenv(name.upper(), raising=False)


def make_settings(**overrides: Any) -> Settings:
    values: dict[str, Any] = {"openai_api_key": "sk-test-key"} | overrides
    return Settings(_env_file=None, **values)


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None, title: str = "") -> bytes:
    document = Document()
    document.core_properties.title = title
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        docx_table = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(docx_table.rows, table):
            for cell, value in zip(row.cells, values):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class FakeEmbedder:
    """Deterministic bag-of-words embedder: texts sharing words get similar vectors."""

    dimension = FAKE_DIMENSION

    def __init__(self) -> None:
        self.calls: list[list[str]] = []

    def embed(self, texts: list[str]) -> list[list[float]]:
        self.calls.append(texts)
        vectors = []
        for text in texts:
            vector = [0.0] * self.dimension
            for word in re.findall(r"\w+", text.lower()):
                vector[zlib.crc32(word.encode()) % self.dimension] += 1.0
            norm = math.sqrt(sum(value * value for value in vector)) or 1.0
            vectors.append([value / norm for value in vector])
        return vectors


class FakeChatModel:
    def __init__(self, answer: str = "Fake answer.") -> None:
        self.answer = answer
        self.prompts: list[tuple[str, str]] = []

    def complete(self, system_prompt: str, user_prompt: str) -> str:
        self.prompts.append((system_prompt, user_prompt))
        return self.answer


class FakePineconeIndex:
    """In-memory stand-in for a Pinecone serverless index data-plane client."""

    def __init__(self, dimension: int = FAKE_DIMENSION, metric: str = "cosine") -> None:
        self.dimension = dimension
        self.metric = metric
        self.namespaces: dict[str, dict[str, dict[str, Any]]] = {}

    def describe_index_stats(self) -> SimpleNamespace:
        return SimpleNamespace(dimension=self.dimension, metric=self.metric)

    def upsert(self, vectors: list[dict[str, Any]], namespace: str, show_progress: bool) -> None:
        records = self.namespaces.setdefault(namespace, {})
        for vector in vectors:
            records[vector["id"]] = vector

    def query(self, vector: list[float], top_k: int, namespace: str, include_metadata: bool) -> SimpleNamespace:
        records = self.namespaces.get(namespace, {}).values()
        scored = sorted(
            (SimpleNamespace(id=r["id"], score=sum(a * b for a, b in zip(vector, r["values"])), metadata=r["metadata"])
             for r in records),
            key=lambda match: match.score,
            reverse=True,
        )
        return SimpleNamespace(matches=scored[:top_k])

    def fetch(self, ids: list[str], namespace: str) -> SimpleNamespace:
        records = self.namespaces.get(namespace, {})
        found = {i: SimpleNamespace(metadata=records[i]["metadata"]) for i in ids if i in records}
        return SimpleNamespace(vectors=found)

    def list(self, prefix: str | None = None, namespace: str = "") -> Iterator[SimpleNamespace]:
        ids = sorted(i for i in self.namespaces.get(namespace, {}) if i.startswith(prefix or ""))
        for start in range(0, len(ids), 100):
            yield SimpleNamespace(vectors=[SimpleNamespace(id=i) for i in ids[start:start + 100]])

    def delete(self, ids: list[str], namespace: str) -> None:
        records = self.namespaces.get(namespace, {})
        for i in ids:
            records.pop(i, None)


@pytest.fixture
def fake_pinecone(monkeypatch: pytest.MonkeyPatch) -> Callable[..., FakePineconeIndex]:
    """Patch the Pinecone SDK client; returns a function that sets the fake index to serve."""
    holder: dict[str, FakePineconeIndex] = {"index": FakePineconeIndex()}

    class FakePinecone:
        def __init__(self, api_key: str) -> None:
            self.api_key = api_key

        def Index(self, name: str) -> FakePineconeIndex:  # noqa: N802 - mirrors the SDK method name
            return holder["index"]

    monkeypatch.setattr("app.vectorstores.pinecone_store.Pinecone", FakePinecone)

    def use(index: FakePineconeIndex | None = None) -> FakePineconeIndex:
        if index is not None:
            holder["index"] = index
        return holder["index"]

    return use
