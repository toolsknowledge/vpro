import pytest
from docx import Document

from app.core.exceptions import InvalidDocumentError
from app.document_loaders.docx_loader import load_docx
from tests.conftest import make_docx


def test_extracts_paragraphs_tables_and_metadata() -> None:
    data = make_docx(["Introduction", "", "Details here."], table=[["Name", "Role"], ["Ada", "Engineer"]],
                     title="Guide")

    loaded = load_docx(data)

    assert loaded.text == "Introduction\nDetails here.\nName | Role\nAda | Engineer"
    assert loaded.title == "Guide"


def test_keeps_document_order_of_tables_and_paragraphs(tmp_path) -> None:
    document = Document()
    document.add_paragraph("Before")
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Cell"
    document.add_paragraph("After")
    path = tmp_path / "ordered.docx"
    document.save(path)

    assert load_docx(path.read_bytes()).text == "Before\nCell\nAfter"


def test_rejects_non_docx_bytes() -> None:
    with pytest.raises(InvalidDocumentError, match="not a readable DOCX"):
        load_docx(b"plain text pretending to be a docx")


def test_rejects_document_without_text() -> None:
    with pytest.raises(InvalidDocumentError, match="no extractable text"):
        load_docx(make_docx(["", "   "]))
