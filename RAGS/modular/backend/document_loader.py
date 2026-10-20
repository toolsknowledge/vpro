from docx import Document

# load_docx() - used to load docx file
def load_docx(file_path):
    """Read paragraphs and tables from a DOCX file."""
    doc = Document(file_path)
    parts = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                parts.append(line)

    return "\n".join(parts)


# used to create chunks
def split_text(text, chunk_size=700, overlap=100):
    """Split text into small overlapping chunks."""
    words = text.split()
    chunks = []
    start = 0

    while start < len(words):
        end = start + chunk_size
        chunk = " ".join(words[start:end]).strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(words):
            break
        start = end - overlap

    return chunks
