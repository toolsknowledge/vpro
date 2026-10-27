"""Text extraction for the supported document formats.

Each document is returned as a list of sections (a PDF page, a spreadsheet
row, or the whole text file) so location metadata can travel with chunks.
"""

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import docx
import openpyxl
from pypdf import PdfReader
from pypdf.errors import PdfReadError

SUPPORTED_EXTENSIONS: dict[str, str] = {
    ".txt": "txt",
    ".md": "md",
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
}


class DocumentLoadError(Exception):
    """Text could not be extracted from a file."""


class UnsupportedFileTypeError(DocumentLoadError):
    pass


class EmptyDocumentError(DocumentLoadError):
    pass


@dataclass
class DocumentSection:
    text: str
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass
class LoadedDocument:
    source: str
    file_type: str
    sections: list[DocumentSection]

    @property
    def full_text(self) -> str:
        return "\n\n".join(section.text for section in self.sections)


def is_supported(path: Path) -> bool:
    return path.suffix.lower() in SUPPORTED_EXTENSIONS


def list_files(directory: Path) -> list[Path]:
    """All regular, non-hidden files in ``directory`` (not recursive), sorted by name."""
    if not directory.is_dir():
        return []
    return sorted(
        (path for path in directory.iterdir() if path.is_file() and not path.name.startswith(".")),
        key=lambda path: path.name.lower(),
    )


def _read_text(path: Path) -> list[DocumentSection]:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    return [DocumentSection(text=text)]


def _read_pdf(path: Path) -> list[DocumentSection]:
    try:
        reader = PdfReader(path)
        return [
            DocumentSection(text=page.extract_text() or "", metadata={"page": number})
            for number, page in enumerate(reader.pages, start=1)
        ]
    except PdfReadError as exc:
        raise DocumentLoadError(f"Could not read PDF: {exc}") from exc


def _read_docx(path: Path) -> list[DocumentSection]:
    try:
        document = docx.Document(str(path))
    except (ValueError, KeyError, OSError) as exc:
        raise DocumentLoadError(f"Could not read DOCX: {exc}") from exc
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [DocumentSection(text="\n".join(blocks))]


def _read_xlsx(path: Path) -> list[DocumentSection]:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (ValueError, KeyError, OSError) as exc:
        raise DocumentLoadError(f"Could not read XLSX: {exc}") from exc
    sections: list[DocumentSection] = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                continue
            columns = [str(value).strip() if value is not None else f"Column {i + 1}" for i, value in enumerate(header)]
            for row_number, row in enumerate(rows, start=2):
                pairs = [
                    f"{columns[i] if i < len(columns) else f'Column {i + 1}'}: {str(value).strip()}"
                    for i, value in enumerate(row)
                    if value is not None and str(value).strip()
                ]
                if pairs:
                    sections.append(
                        DocumentSection(text="\n".join(pairs), metadata={"sheet": sheet.title, "row": row_number})
                    )
    finally:
        workbook.close()
    return sections


_READERS = {
    "txt": _read_text,
    "md": _read_text,
    "pdf": _read_pdf,
    "docx": _read_docx,
    "xlsx": _read_xlsx,
}


def load_document(path: Path) -> LoadedDocument:
    file_type = SUPPORTED_EXTENSIONS.get(path.suffix.lower())
    if file_type is None:
        raise UnsupportedFileTypeError(f"Unsupported file type '{path.suffix or '(none)'}'")
    sections = [
        DocumentSection(text=section.text.strip(), metadata=section.metadata)
        for section in _READERS[file_type](path)
        if section.text and section.text.strip()
    ]
    if not sections:
        raise EmptyDocumentError("No extractable text found")
    return LoadedDocument(source=path.name, file_type=file_type, sections=sections)
