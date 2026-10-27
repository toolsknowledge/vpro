"""Generate the sample knowledge base and sample questions.

Usage (from the project root):
    .venv/bin/python backend/scripts/create_sample_files.py           # create missing files only
    .venv/bin/python backend/scripts/create_sample_files.py --force   # recreate every file

Content is fixed in this script and document timestamps are pinned, so the
extracted text is identical on every run (the indexer hashes extracted text,
so regenerated files are not re-embedded unless their content changes).
"""

import argparse
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))

import docx  # noqa: E402
import openpyxl  # noqa: E402
from fpdf import FPDF  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

from config import SAMPLE_DOCUMENTS_DIR  # noqa: E402

SAMPLE_QUESTIONS_FILE = BACKEND_DIR / "data" / "sample_questions.json"
FIXED_TIMESTAMP = datetime(2025, 1, 1, tzinfo=timezone.utc)

RAG_FUNDAMENTALS = """RAG FUNDAMENTALS - VPRO Skills Course Notes

1. What is Retrieval-Augmented Generation (RAG)?
Retrieval-Augmented Generation (RAG) is a technique that combines a large language model (LLM) with an external knowledge source. Instead of relying only on what the model memorised during training, a RAG system first retrieves relevant passages from a document collection and then gives those passages to the LLM as context. The LLM writes its answer using that context. RAG lets an assistant answer questions about private, recent, or domain-specific information such as course notes, company policies, or product manuals.

2. The two phases of RAG
A RAG system has an offline indexing phase and an online retrieval-and-generation phase.

Indexing phase: documents are loaded, their text is extracted, the text is split into chunks, each chunk is converted into an embedding vector, and the vectors are stored in a vector database together with the original text and metadata such as the file name.

Retrieval and generation phase: when a user asks a question, the question is converted into an embedding with the same embedding model, the vector database returns the chunks whose embeddings are most similar to the question, and the LLM generates an answer grounded in those chunks.

3. Embeddings
An embedding is a list of numbers (a vector) that represents the meaning of a piece of text. Texts with similar meaning have vectors that point in similar directions. Similarity is usually measured with cosine similarity, where 1.0 means identical direction and values near 0 mean unrelated. In this course we use the OpenAI model text-embedding-3-small, which produces 1536-dimensional vectors. The same embedding model must be used for documents and for queries, otherwise the vectors are not comparable.

4. Chunks and chunking
Documents are usually too long to embed or to send to an LLM in one piece, so they are split into chunks. A typical chunk is a few hundred to a thousand characters. Chunk overlap repeats a small amount of text between neighbouring chunks so that a sentence cut at a boundary still appears whole in at least one chunk. Chunk size is a trade-off: small chunks are precise but may lose context, large chunks keep context but dilute similarity scores and use more of the LLM context window. Good chunkers split at natural boundaries such as paragraphs and sentences.

5. Vector databases
A vector database stores embeddings and finds the nearest vectors to a query vector quickly. ChromaDB is an open-source vector database that can run locally and persist its data to disk. Each record holds an id, the embedding, the chunk text, and metadata. Metadata such as source file name, page number, or chunk index makes it possible to show citations to users.

6. Advantages of RAG
- Answers can use private and up-to-date information without retraining the model.
- Answers can cite their sources, which improves trust and makes verification easy.
- Hallucination is reduced because the model is instructed to answer from supplied context.
- Updating knowledge is cheap: re-index the changed documents instead of fine-tuning.
- Access control can be applied at retrieval time.

7. Limitations of RAG
- Answer quality depends on retrieval quality: if the right chunk is not retrieved, the answer will be incomplete or wrong.
- Poor chunking can split important information across chunks.
- Every query that retrieves costs an embedding call and adds latency.
- Retrieved context can be irrelevant; without a relevance check the LLM may be distracted or may invent connections.
- A fixed pipeline retrieves even for greetings and general questions, wasting time and money. Adaptive RAG addresses this.
"""

ADAPTIVE_RAG_GUIDE = """# Adaptive RAG Guide

## Definition

**Adaptive RAG** is a Retrieval-Augmented Generation design in which the system decides, for every user query,
*whether* and *how* to retrieve before answering. A standard (naive) RAG pipeline always retrieves, even for
"Hello". Adaptive RAG first classifies the query and then follows the cheapest strategy that can answer it well.
This reduces cost and latency and avoids forcing irrelevant document context into simple conversations.

## The three routes used in this course

| Route | When it is used | What happens |
|-------|-----------------|--------------|
| DIRECT | Greetings, small talk, general-knowledge questions not tied to the course documents | The LLM answers directly. No vector search is performed and no sources are shown. |
| RETRIEVE | Questions about the course notes or indexed documents | The query is embedded, ChromaDB returns the top-k chunks, a relevance check filters them, and the LLM answers only from the remaining context with citations. |
| CLARIFY | Ambiguous or incomplete questions such as "Explain it" with no earlier topic | The assistant asks a short follow-up question. No retrieval is performed. |

## Query classification

The router is an LLM call with a strict output schema. It receives the question, a short window of conversation
history, and the list of indexed documents, and must return JSON with three fields:

- `route` - one of DIRECT, RETRIEVE, CLARIFY
- `reason` - one sentence explaining the decision (shown to the learner)
- `clarifying_question` - required when the route is CLARIFY

The output is validated with a Pydantic model. If it is malformed, the system uses a documented fallback route
instead of crashing. The router output is treated purely as data: it never executes code or tools.

Conversation history matters. "What about the second one?" is ambiguous on its own, but if the previous answer
listed Hybrid RAG and Advanced RAG, the router can resolve "the second one" to Advanced RAG and retrieve.

## Retrieval relevance checks

Retrieving is not enough - the retrieved chunks must actually be relevant. Each chunk returned by ChromaDB has a
cosine similarity score between the question and the chunk. Chunks below a configurable **relevance threshold**
(0.35 by default in this project) are discarded. If no chunk passes the threshold, the assistant says that the
indexed documents do not contain enough information instead of inventing an answer.

## Example routing scenarios

1. "Hi, how are you?" -> DIRECT. A greeting needs no documents.
2. "What is RAG in simple words?" -> usually DIRECT. It is a general question that the LLM can answer.
3. "Explain Adaptive RAG from the course notes." -> RETRIEVE. The learner explicitly asks for the notes.
4. "What is the difference between Hybrid RAG and Advanced RAG in our documents?" -> RETRIEVE.
5. "Explain it." as the first message -> CLARIFY. There is no topic to explain.
6. "Compare these two." after a message that named two techniques -> RETRIEVE with the references resolved.
7. "What is the capital of France according to our notes?" -> RETRIEVE, but the relevance check finds nothing,
   so the assistant reports that the documents do not cover it.

## Benefits and trade-offs

- Benefit: fewer embedding and vector-search calls, lower latency for simple messages.
- Benefit: clarification prevents confident answers to questions the system did not understand.
- Trade-off: the classifier is an extra LLM call and can misroute borderline questions, so routing decisions
  should be visible to users and the relevance check acts as a safety net.
"""

HYBRID_VS_ADVANCED = [
    ("heading", "Hybrid RAG vs Advanced RAG"),
    ("paragraph", "VPRO Skills course notes comparing two families of RAG improvements."),
    ("heading2", "Hybrid RAG"),
    ("paragraph", "Hybrid RAG combines more than one retrieval method, most commonly sparse keyword retrieval and "
                  "dense vector retrieval, and merges their results. The goal is to get the precision of exact "
                  "keyword matching together with the semantic understanding of embeddings."),
    ("heading2", "Sparse retrieval"),
    ("paragraph", "Sparse retrieval represents text as word-count style vectors in which most values are zero. "
                  "BM25 and TF-IDF are classic sparse methods. They excel at exact terms such as product codes, "
                  "error messages, names, and acronyms, but they miss synonyms: a search for 'car' does not "
                  "match 'automobile'."),
    ("heading2", "Dense retrieval"),
    ("paragraph", "Dense retrieval uses embedding models to map text into dense vectors where nearby vectors have "
                  "similar meaning. It finds paraphrases and related concepts, but it can miss rare exact tokens "
                  "and is harder to explain. ChromaDB with OpenAI text-embedding-3-small is a dense retriever."),
    ("paragraph", "Hybrid search runs both retrievers and fuses the ranked lists, often with Reciprocal Rank "
                  "Fusion (RRF), which scores each document by summing 1 / (k + rank) across the lists."),
    ("heading2", "Advanced RAG"),
    ("paragraph", "Advanced RAG improves the naive retrieve-then-generate pipeline with extra steps before and "
                  "after retrieval. Pre-retrieval steps improve the query; post-retrieval steps improve the "
                  "context that reaches the LLM."),
    ("heading2", "Query rewriting"),
    ("paragraph", "Query rewriting uses an LLM to turn a vague or conversational question into a precise search "
                  "query, for example by resolving pronouns from chat history, expanding abbreviations, or "
                  "generating several alternative queries (multi-query retrieval)."),
    ("heading2", "Reranking"),
    ("paragraph", "Reranking retrieves a larger candidate set (for example the top 20 chunks) and then re-scores "
                  "each candidate against the question with a more accurate but slower model, such as a "
                  "cross-encoder, keeping only the best few. This raises precision of the final context."),
    ("heading2", "Key differences"),
    ("table", [
        ("Aspect", "Hybrid RAG", "Advanced RAG"),
        ("Main idea", "Combine sparse and dense retrieval", "Add pre- and post-retrieval optimisation steps"),
        ("Typical techniques", "BM25 + embeddings, Reciprocal Rank Fusion", "Query rewriting, reranking, compression"),
        ("Best for", "Mixed queries with exact terms and concepts", "Complex or poorly phrased questions"),
        ("Extra cost", "A second index and a fusion step", "Additional LLM or reranker calls"),
    ]),
    ("heading2", "Practical examples"),
    ("paragraph", "Example 1 - Hybrid RAG: a support bot must match the exact error code 'ERR-4031' (sparse) and "
                  "also understand 'payment keeps failing' (dense). Hybrid search covers both."),
    ("paragraph", "Example 2 - Advanced RAG: a learner asks 'how does it pick chunks?' after discussing ChromaDB. "
                  "Query rewriting turns this into 'How does ChromaDB select the most similar chunks for a "
                  "query?', and a reranker then keeps the three most relevant passages."),
    ("paragraph", "The two approaches are complementary: a production system can use hybrid retrieval inside an "
                  "advanced pipeline, and add adaptive routing on top to decide when retrieval is needed at all."),
]

FAQ = [
    ("What does RAG stand for?",
     "RAG stands for Retrieval-Augmented Generation. A RAG system retrieves relevant passages from a knowledge base "
     "and passes them to a large language model so that the answer is grounded in those passages."),
    ("Why not just ask the LLM directly?",
     "An LLM only knows what was in its training data. It does not know your private course notes or recent "
     "changes, and it may hallucinate. RAG supplies the needed facts at question time and allows citations."),
    ("What is Adaptive RAG?",
     "Adaptive RAG classifies each query and chooses a strategy: answer directly (DIRECT), retrieve from documents "
     "(RETRIEVE), or ask a follow-up question (CLARIFY). It avoids unnecessary retrieval for simple messages."),
    ("When does the assistant ask a clarifying question?",
     "When a question is incomplete or ambiguous and the conversation history does not resolve it, for example "
     "'Explain it' as the very first message. Clarifying avoids answering the wrong question."),
    ("What is Hybrid RAG?",
     "Hybrid RAG combines sparse keyword retrieval such as BM25 with dense embedding retrieval and fuses the "
     "results, often using Reciprocal Rank Fusion. It handles both exact terms and semantic matches."),
    ("What is ChromaDB?",
     "ChromaDB is an open-source vector database. It stores embeddings, chunk text, and metadata, supports "
     "similarity search, and can persist data to a local directory so the index survives restarts."),
    ("Where does this project store its ChromaDB data?",
     "In a persistent local directory configured by CHROMA_PERSIST_DIRECTORY, by default backend/storage/chroma. "
     "The collection name is configured by CHROMA_COLLECTION_NAME."),
    ("What are embeddings?",
     "Embeddings are numeric vectors that capture the meaning of text. Similar texts have similar vectors, so "
     "comparing vectors with cosine similarity finds passages related to a question."),
    ("Which OpenAI models does the course project use?",
     "By default it uses gpt-4.1-mini for routing and answer generation and text-embedding-3-small for "
     "embeddings. Both are configurable with OPENAI_CHAT_MODEL and OPENAI_EMBEDDING_MODEL."),
    ("What is the relevance threshold?",
     "It is the minimum cosine similarity a retrieved chunk must have to be used as context. Chunks below the "
     "threshold are ignored; if none pass, the assistant says the documents do not contain the answer."),
    ("Does using the OpenAI API cost money?",
     "Yes. Embedding documents and generating answers are billed by OpenAI per token. Indexing the small sample "
     "documents costs very little, and unchanged files are not re-embedded."),
    ("What is chunk overlap?",
     "Chunk overlap repeats a small number of characters between consecutive chunks so that information near a "
     "chunk boundary is not lost. This project uses 120 characters of overlap with 800-character chunks."),
]

CONCEPTS = [
    ("Retrieval-Augmented Generation", "Combining document retrieval with LLM generation so answers are grounded in "
     "retrieved text.", "A course assistant answering from uploaded lecture notes."),
    ("Embedding", "A numeric vector representing the meaning of text.",
     "text-embedding-3-small turns a paragraph into a 1536-number vector."),
    ("Chunk", "A small piece of a document that is embedded and retrieved independently.",
     "An 800-character section of a PDF page."),
    ("Chunk overlap", "Text repeated between neighbouring chunks to preserve context at boundaries.",
     "The last 120 characters of chunk 3 also start chunk 4."),
    ("Vector database", "A database that stores embeddings and finds the nearest vectors to a query.",
     "ChromaDB persisting vectors in backend/storage/chroma."),
    ("Cosine similarity", "A measure of how closely two vectors point in the same direction, from -1 to 1.",
     "A score of 0.82 between a question and a highly relevant chunk."),
    ("Top-k retrieval", "Returning the k most similar chunks for a query.",
     "TOP_K=5 returns the five closest chunks."),
    ("Relevance threshold", "Minimum similarity a chunk needs to be used as answer context.",
     "Chunks scoring below 0.35 are discarded."),
    ("Adaptive RAG", "RAG that decides per query whether to answer directly, retrieve, or clarify.",
     "Answering 'Hi' without searching the documents."),
    ("Query classification", "Using a model to assign a route to each user query.",
     "An LLM returns {\"route\": \"RETRIEVE\", \"reason\": \"asks about course notes\"}."),
    ("Hybrid RAG", "Combining sparse keyword search with dense vector search.",
     "BM25 plus embeddings merged with Reciprocal Rank Fusion."),
    ("Advanced RAG", "RAG with pre-retrieval and post-retrieval optimisation steps.",
     "Query rewriting before search and reranking after search."),
    ("Reranking", "Re-scoring retrieved candidates with a more accurate model and keeping the best.",
     "A cross-encoder selects the top 3 of 20 retrieved chunks."),
    ("Hallucination", "A fluent but unsupported or false statement generated by an LLM.",
     "Inventing a course deadline that appears in no document."),
]

SAMPLE_QUESTIONS = {
    "note": "Expected routes are what the classifier is designed to choose. The route actually taken is decided "
            "per question by the LLM and shown on every answer; borderline questions may route differently.",
    "DIRECT": [
        "Hi, how are you?",
        "Hello",
        "What is RAG in simple words?",
        "Thanks for the help!",
    ],
    "RETRIEVE": [
        "Explain Adaptive RAG from the course notes.",
        "What is the difference between Hybrid RAG and Advanced RAG in our documents?",
        "What is ChromaDB according to the sample documents?",
        "What are the limitations of RAG described in the notes?",
        "What relevance threshold does the course project use?",
        "Which concepts are defined in the concepts spreadsheet?",
    ],
    "CLARIFY": [
        "Explain it.",
        "What about the second one?",
        "Compare these two.",
    ],
    "BORDERLINE": [
        "What are embeddings?",
        "How does reranking work?",
    ],
    "OUT_OF_SCOPE": [
        "What is the capital of France according to our course notes?",
        "What does our documentation say about Kubernetes pod autoscaling?",
    ],
}


def write_txt(path: Path) -> None:
    path.write_text(RAG_FUNDAMENTALS, encoding="utf-8")


def write_md(path: Path) -> None:
    path.write_text(ADAPTIVE_RAG_GUIDE, encoding="utf-8")


def write_docx(path: Path) -> None:
    document = docx.Document()
    for kind, content in HYBRID_VS_ADVANCED:
        if kind == "heading":
            document.add_heading(content, level=1)
        elif kind == "heading2":
            document.add_heading(content, level=2)
        elif kind == "paragraph":
            document.add_paragraph(content)
        elif kind == "table":
            table = document.add_table(rows=0, cols=len(content[0]))
            table.style = "Table Grid"
            for row in content:
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = value
    properties = document.core_properties
    properties.author = "VPRO Skills EduTech"
    properties.title = "Hybrid RAG vs Advanced RAG"
    properties.created = FIXED_TIMESTAMP
    properties.modified = FIXED_TIMESTAMP
    properties.last_modified_by = "VPRO Skills EduTech"
    document.save(str(path))


def write_pdf(path: Path) -> None:
    pdf = FPDF(format="A4")
    pdf.set_creation_date(FIXED_TIMESTAMP)
    pdf.set_title("RAG Course FAQ")
    pdf.set_author("VPRO Skills EduTech")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 12, "RAG Course FAQ - VPRO Skills", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 8, "Frequently asked questions about RAG, Adaptive RAG, Hybrid RAG, ChromaDB, embeddings and OpenAI.",
             new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    for number, (question, answer) in enumerate(FAQ, start=1):
        pdf.set_font("Helvetica", "B", 12)
        pdf.multi_cell(0, 7, f"Q{number}. {question}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 6, f"A: {answer}", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)
    pdf.output(str(path))


def write_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Concepts"
    sheet.append(["Concept", "Definition", "Example"])
    for row in CONCEPTS:
        sheet.append(list(row))
    header_fill = PatternFill("solid", fgColor="F7931E")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for column, width in zip("ABC", (28, 70, 60)):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.properties.creator = "VPRO Skills EduTech"
    workbook.properties.created = FIXED_TIMESTAMP.replace(tzinfo=None)
    workbook.properties.modified = FIXED_TIMESTAMP.replace(tzinfo=None)
    workbook.save(path)


def write_questions(path: Path) -> None:
    path.write_text(json.dumps(SAMPLE_QUESTIONS, indent=2) + "\n", encoding="utf-8")


GENERATORS = {
    SAMPLE_DOCUMENTS_DIR / "rag_fundamentals.txt": write_txt,
    SAMPLE_DOCUMENTS_DIR / "adaptive_rag_guide.md": write_md,
    SAMPLE_DOCUMENTS_DIR / "hybrid_vs_advanced_rag.docx": write_docx,
    SAMPLE_DOCUMENTS_DIR / "rag_course_faq.pdf": write_pdf,
    SAMPLE_DOCUMENTS_DIR / "rag_concepts.xlsx": write_xlsx,
    SAMPLE_QUESTIONS_FILE: write_questions,
}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--force", action="store_true", help="overwrite existing files")
    args = parser.parse_args()

    SAMPLE_DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)
    for path, generate in GENERATORS.items():
        if path.exists() and not args.force:
            print(f"  exists   {path.relative_to(BACKEND_DIR)}")
            continue
        generate(path)
        print(f"  created  {path.relative_to(BACKEND_DIR)} ({path.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
