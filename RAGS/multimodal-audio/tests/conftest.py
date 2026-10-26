"""Shared test fixtures: sample files built in memory and a fake OpenAI client.

No test in this folder (except tests/integration) calls the real OpenAI API.
"""

from __future__ import annotations

import hashlib
import io
import math
import re
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from backend.document_registry import DocumentRegistry
from backend.embeddings import EmbeddingService
from backend.indexing_pipeline import IndexingPipeline
from backend.vector_store import VectorStore

FAKE_DIM = 64


def fake_vector(text: str, dim: int = FAKE_DIM) -> list[float]:
    """Deterministic bag-of-words vector: texts sharing words get similar vectors."""
    vec = [0.0] * dim
    for word in re.findall(r"[a-z0-9]+", text.lower()):
        bucket = int(hashlib.md5(word.encode()).hexdigest(), 16) % dim
        vec[bucket] += 1.0
    norm = math.sqrt(sum(v * v for v in vec)) or 1.0
    return [v / norm for v in vec]


class FakeEmbeddings:
    def __init__(self, dim=FAKE_DIM, fail_times: int = 0, error: Exception | None = None):
        self.dim, self.fail_times, self.error, self.calls = dim, fail_times, error, 0

    def create(self, model, input):
        self.calls += 1
        if self.fail_times > 0:
            self.fail_times -= 1
            raise self.error
        data = [SimpleNamespace(index=i, embedding=fake_vector(t, self.dim)) for i, t in enumerate(input)]
        return SimpleNamespace(data=list(reversed(data)))  # reversed: the code must sort by index


class FakeChat:
    def __init__(self, reply: str = "Answer [S1]."):
        self.reply, self.last_kwargs = reply, None
        self.completions = self

    def create(self, **kwargs):
        self.last_kwargs = kwargs
        message = SimpleNamespace(content=self.reply)
        usage = SimpleNamespace(prompt_tokens=10, completion_tokens=5)
        return SimpleNamespace(choices=[SimpleNamespace(message=message)], usage=usage)


class FakeClient:
    def __init__(self, reply: str = "Answer [S1].", **embedding_kwargs):
        self.embeddings = FakeEmbeddings(**embedding_kwargs)
        self.chat = FakeChat(reply)


# --------------------------------------------------------------- file builders
def make_pdf_bytes(pages: list[str]) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    for text in pages:
        y = 800
        for line in text.split("\n"):
            if line:
                pdf.drawString(50, y, line)
            y -= 16
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_xlsx_bytes(sheets: dict[str, list[tuple]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def make_docx_bytes(blocks: list[tuple[str, str]], table: list[tuple] | None = None) -> bytes:
    """blocks = [("heading", text) | ("para", text)]"""
    doc = Document()
    for kind, text in blocks:
        if kind == "heading":
            doc.add_heading(text, 1)
        else:
            doc.add_paragraph(text)
    if table:
        t = doc.add_table(rows=0, cols=len(table[0]))
        for row in table:
            cells = t.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = str(value)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def course_xlsx() -> bytes:
    return make_xlsx_bytes(
        {
            "Course Catalog": [
                ("Course", "Duration", "Fee"),
                ("Python Full Stack", "16 weeks", 30000),
                ("Data Science", "20 weeks", None),
            ],
            "Batches": [("Course", "Start"), ("Python Full Stack", "2026-11-10")],
        }
    )


@pytest.fixture
def services_factory(tmp_path):
    """Build a real registry + real ChromaDB (in tmp_path) + fake OpenAI embeddings."""

    def build(client=None, model="fake-embedding-model", collection="test_collection"):
        client = client or FakeClient()
        registry = DocumentRegistry(tmp_path / "registry.json", tmp_path / "traces", tmp_path / "uploads")
        store = VectorStore(tmp_path / "chroma", collection)
        embedder = EmbeddingService(client, model, batch_size=4, max_attempts=3)
        pipeline = IndexingPipeline(registry, store, embedder, max_upload_bytes=5 * 1024 * 1024)
        return SimpleNamespace(client=client, registry=registry, store=store, embedder=embedder, pipeline=pipeline)

    return build


@pytest.fixture(autouse=True)
def no_sleep(monkeypatch):
    """Retries back off with time.sleep; skip the waiting in tests."""
    monkeypatch.setattr("backend.openai_client.time.sleep", lambda seconds: None)


# =================================================================== media
# Fakes for the audio/video pipeline. Real FFmpeg is used to build tiny test
# media files; tests that need it are skipped when FFmpeg is not installed.
import json
import shutil
import subprocess

from backend.frame_analyzer import FrameAnalyzer
from backend.media_pipeline import MediaIndexingPipeline
from backend.media_registry import MediaRegistry
from backend.media_utils import FFmpegTools
from backend.media_vector_store import MediaStores
from backend.transcription_service import TranscriptionService

HAS_FFMPEG = bool(shutil.which("ffmpeg") and shutil.which("ffprobe"))
needs_ffmpeg = pytest.mark.skipif(not HAS_FFMPEG, reason="FFmpeg is not installed")

PIECE_TOPICS = [
    "Welcome. Today we study Python list comprehensions such as squares = [n * n for n in range(10)].",
    "Embeddings turn text into vectors; text-embedding-3-small has 1536 dimensions.",
    "Chunk overlap of 120 tokens keeps sentences intact. ChromaDB stores the vectors on disk.",
    "RAG evaluation checks recall at K, faithfulness and answer relevance.",
]


class FakeTranscriptions:
    """Mimics client.audio.transcriptions.create for json / verbose_json / diarized_json."""

    def __init__(self, fail_pieces: set[int] | None = None, error: Exception | None = None):
        self.calls: list[dict] = []
        self.fail_pieces = set(fail_pieces or ())
        self.error = error

    def create(self, file, model, response_format="json", **kwargs):
        name = getattr(file, "name", "")
        piece = int(re.search(r"piece_(\d+)", name).group(1)) if "piece_" in name else 0
        self.calls.append({"model": model, "response_format": response_format, "piece": piece, **kwargs})
        if piece in self.fail_pieces:
            raise self.error
        text = PIECE_TOPICS[piece % len(PIECE_TOPICS)]
        if response_format == "verbose_json":
            half = len(text) // 2
            return SimpleNamespace(
                text=text, language="english",
                segments=[SimpleNamespace(start=0.0, end=3.0, text=text[:half]), SimpleNamespace(start=3.0, end=6.5, text=text[half:])],
            )
        if response_format == "diarized_json":
            return SimpleNamespace(text=text, segments=[
                SimpleNamespace(id="seg_0", start=0.2, end=4.0, text=text, speaker="A", type="transcript.text.segment"),
            ])
        return SimpleNamespace(text=text)


class FakeMediaChat:
    """Answers vision requests with JSON descriptions and RAG requests with a fixed reply."""

    def __init__(self, reply: str = "The instructor explained list comprehensions [S1].", vision_error: Exception | None = None,
                 vision_fail_times: int = 0):
        self.reply = reply
        self.vision_error = vision_error
        self.vision_fail_times = vision_fail_times
        self.vision_calls = 0
        self.rag_calls: list[dict] = []
        self.completions = self

    def create(self, **kwargs):
        content = kwargs["messages"][-1]["content"]
        if isinstance(content, list):  # vision request: [text, image_url]
            self.vision_calls += 1
            if self.vision_error and (self.vision_fail_times < 0 or self.vision_fail_times > 0):
                self.vision_fail_times -= 1
                raise self.vision_error
            label = content[0]["text"]
            payload = {
                "summary": f"A slide about ChromaDB vector storage. {label}",
                "slide_title": "Storing Vectors in ChromaDB",
                "visible_text": "PersistentClient keeps data on disk",
                "code": "client = chromadb.PersistentClient(path='data/chroma_db')",
                "diagrams_tables_charts": "",
                "people_and_actions": "",
                "uncertain": "",
            }
            text = json.dumps(payload)
        else:
            self.rag_calls.append(kwargs)
            text = self.reply
        message = SimpleNamespace(content=text)
        return SimpleNamespace(choices=[SimpleNamespace(message=message)],
                               usage=SimpleNamespace(prompt_tokens=10, completion_tokens=5))


class FakeMediaClient:
    def __init__(self, transcriptions: FakeTranscriptions | None = None, chat: FakeMediaChat | None = None, **embedding_kwargs):
        self.embeddings = FakeEmbeddings(**embedding_kwargs)
        self.audio = SimpleNamespace(transcriptions=transcriptions or FakeTranscriptions())
        self.chat = chat or FakeMediaChat()


def run_ffmpeg(*args: str) -> None:
    subprocess.run(["ffmpeg", "-hide_banner", "-loglevel", "error", "-nostdin", "-y", *args], check=True)


def make_tone(path, seconds: float = 70, gap_every: float = 10, gap: float = 1.5):
    """A 440 Hz tone with a silent gap every `gap_every` seconds (for silence detection)."""
    run_ffmpeg("-f", "lavfi", "-i", f"sine=frequency=440:duration={seconds}",
               "-af", f"volume='if(lt(mod(t,{gap_every}),{gap_every - gap}),1,0)':eval=frame", str(path))
    return path


def make_video(path, colors=("red", "red", "blue", "green"), seg_seconds: float = 10, audio: bool = True):
    """A video made of solid-colour blocks (identical blocks = duplicate frames), optionally with a tone."""
    inputs, filters = [], []
    for i, color in enumerate(colors):
        inputs += ["-f", "lavfi", "-i", f"color=c={color}:s=320x240:r=5:d={seg_seconds}"]
        filters.append(f"[{i}:v]")
    graph = "".join(filters) + f"concat=n={len(colors)}:v=1:a=0[v]"
    args = inputs
    if audio:
        args += ["-f", "lavfi", "-i", f"sine=frequency=300:duration={seg_seconds * len(colors)}"]
    args += ["-filter_complex", graph, "-map", "[v]"]
    if audio:
        args += ["-map", f"{len(colors)}:a", "-c:a", "aac"]
    args += ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-shortest", str(path)]
    run_ffmpeg(*args)
    return path


@pytest.fixture
def media_factory(tmp_path):
    """Real registry + real ChromaDB + real FFmpeg, with fake OpenAI clients."""

    def build(client=None, transcription_model="gpt-4o-mini-transcribe", embedding_model="fake-embedding-model",
              audio_chunk_seconds=300, fallback_seconds=20, data_dir=None):
        client = client or FakeMediaClient()
        data = data_dir or (tmp_path / "data")
        registry = MediaRegistry(data / "media_registry.json", data / "transcripts", data / "uploads" / "media",
                                 data / "extracted_audio", data / "video_frames")
        stores = MediaStores(data / "chroma_db", "t_audio", "t_video_transcripts", "t_video_frames")
        embedder = EmbeddingService(client, embedding_model, batch_size=8)
        transcriber = TranscriptionService(client, transcription_model)
        analyzer = FrameAnalyzer(client, "fake-vision-model")
        pipeline = MediaIndexingPipeline(
            registry, stores, FFmpegTools(timeout=120), embedder, transcriber, analyzer, data,
            max_upload_bytes=50 * 1024 * 1024, max_duration_seconds=3600,
            audio_chunk_seconds=audio_chunk_seconds, timestamp_fallback_seconds=fallback_seconds,
        )
        return SimpleNamespace(client=client, registry=registry, stores=stores, embedder=embedder,
                               transcriber=transcriber, analyzer=analyzer, pipeline=pipeline, data=data)

    return build
