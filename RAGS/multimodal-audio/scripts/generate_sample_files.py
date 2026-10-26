"""Create four small demo documents in sample_files/ for classroom testing.

Run:  python scripts/generate_sample_files.py

The content is fictional demo material for VPro Skills EduTech. It is only
there so you have one PDF, one Excel, one Word and one TXT file to upload.
If you already have the official sample files, copy those into sample_files/
instead; this script does not overwrite existing files unless --force is given.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

OUT = Path(__file__).resolve().parent.parent / "sample_files"

PDF_PAGES = [
    (
        "RAG Architecture Notes: Page 1",
        [
            "Retrieval-Augmented Generation (RAG) combines a search step with a language model. Instead of "
            "relying only on what the model memorised during training, RAG retrieves relevant passages from "
            "your own documents and supplies them as context to the language model.",
            "A RAG system has two pipelines. The indexing pipeline prepares documents ahead of time. The "
            "question-answering pipeline runs every time a user asks a question.",
        ],
    ),
    (
        "The Indexing Pipeline: Page 2",
        [
            "Step 1, Upload: the user provides PDF, Excel, Word or text files. Step 2, Extract: text is read "
            "from each file together with metadata such as page numbers and sheet names. Step 3, Clean: "
            "whitespace is normalised without removing meaningful content.",
            "Step 4, Chunk: long text is split into smaller overlapping pieces called chunks, typically a few "
            "hundred tokens each, because embedding models and prompts have size limits and smaller pieces "
            "give more precise retrieval. Step 5, Embed: each chunk is converted into an embedding, a numeric "
            "vector that represents its meaning. Step 6, Store: chunks, vectors and metadata are saved in a "
            "vector database such as ChromaDB.",
        ],
    ),
    (
        "Chunking versus Embeddings: Page 3",
        [
            "Chunking and embeddings are different steps. Chunking decides how the text is divided; it "
            "produces pieces of text. Embedding converts each piece of text into numbers; it produces vectors. "
            "Chunking happens first, and embeddings are computed for every chunk.",
            "At question time, the question is embedded with the same model, the closest chunk vectors are "
            "retrieved from the vector database, and the chunks are passed to the language model, which writes "
            "an answer that cites its sources.",
        ],
    ),
]

COURSES = [
    ("Course", "Duration", "Fee (INR)", "Mode", "Level"),
    ("Python Full Stack", "16 weeks", 30000, "Online + Classroom", "Beginner"),
    ("Data Science with Python", "20 weeks", 45000, "Online", "Intermediate"),
    ("Generative AI and RAG", "8 weeks", 25000, "Online", "Intermediate"),
    ("Java Full Stack", "18 weeks", 32000, "Classroom", "Beginner"),
    ("DevOps and Cloud", "12 weeks", 28000, "Online", "Intermediate"),
]

BATCHES = [
    ("Course", "Batch start", "Timing", "Seats"),
    ("Python Full Stack", "2026-11-10", "7:00 AM to 9:00 AM", 30),
    ("Generative AI and RAG", "2026-11-17", "8:00 PM to 10:00 PM", 25),
    ("Data Science with Python", "2026-12-01", "6:30 PM to 8:30 PM", 30),
]


def make_pdf(path: Path) -> None:
    styles = getSampleStyleSheet()
    story = []
    for i, (title, paragraphs) in enumerate(PDF_PAGES):
        story.append(Paragraph(title, styles["Heading1"]))
        for text in paragraphs:
            story.extend([Paragraph(text, styles["BodyText"]), Spacer(1, 10)])
        if i < len(PDF_PAGES) - 1:
            story.append(PageBreak())
    SimpleDocTemplate(str(path), pagesize=A4, title="RAG Architecture Notes").build(story)


def make_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Course Catalog"
    for row in COURSES:
        ws.append(row)
    ws2 = wb.create_sheet("Upcoming Batches")
    for row in BATCHES:
        ws2.append(row)
    wb.save(path)


def make_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("VPro Skills EduTech: Training Policy", 0)
    doc.add_heading("Attendance", 1)
    doc.add_paragraph(
        "Learners must attend at least 80% of live sessions to be eligible for the course certificate. "
        "Recorded sessions are available for 6 months after the batch ends."
    )
    doc.add_heading("Project assessment", 1)
    doc.add_paragraph(
        "Every course ends with a capstone project. Projects are assessed by a mentor on functionality (40%), "
        "code quality (30%), documentation (15%) and the final presentation (15%). A minimum score of 60% "
        "is required to pass. Learners who do not pass may resubmit once within 30 days."
    )
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Criterion", "Weight"
    for criterion, weight in (("Functionality", "40%"), ("Code quality", "30%"), ("Documentation", "15%"), ("Presentation", "15%")):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = criterion, weight
    doc.add_heading("Refund policy", 1)
    doc.add_paragraph(
        "A full refund is available within 7 days of the batch start date. After 7 days, fees are not refundable, "
        "but learners may transfer to a later batch once."
    )
    doc.save(path)


TXT = """VPro Skills EduTech: Course Overview

Introduction:
VPro Skills EduTech offers job-oriented training programmes with live classes, mentor support and
hands-on projects. This overview describes the Generative AI and RAG course.

Course modules:
1. Python refresher and working with APIs
2. Large language models and prompt engineering
3. Document extraction: PDF, Word, Excel and text files
4. Chunking strategies and token limits
5. Embeddings and vector databases (ChromaDB)
6. Building a RAG application with Streamlit
7. Evaluation, grounding and citations

Learning outcomes:
By the end of the course, learners can build a document question-answering application that retrieves
relevant chunks from a vector database and generates grounded answers with source citations.

Support:
Mentor doubt-clearing sessions are held every Saturday from 10:00 AM to 12:00 PM.
"""


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--force", action="store_true", help="overwrite existing files")
    args = parser.parse_args()
    OUT.mkdir(exist_ok=True)
    builders = {
        "sample_rag_architecture_notes.pdf": make_pdf,
        "sample_course_catalog.xlsx": make_xlsx,
        "sample_training_policy.docx": make_docx,
        "sample_course_overview.txt": lambda p: p.write_text(TXT, encoding="utf-8"),
    }
    for name, build in builders.items():
        path = OUT / name
        if path.exists() and not args.force:
            print(f"skip   {path} (exists)")
            continue
        build(path)
        print(f"wrote  {path}")


if __name__ == "__main__":
    main()
